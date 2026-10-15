"""
Generate a single consolidated BOM across analog + frontpanel + ss_pa.

Reuses the netlist parser, DNP collector, sheet mapper, and category
rules from gen_board_boms.py so the two scripts stay in sync.

For each category tab:
- Rows grouped by Value
- Qty = total across all boards (excluding DNP)
- References column shows board-prefixed refs
  (e.g. "analog:R12, ss_pa:R9")
- Boards column lists which boards contribute to the row
- Footprint shows the single footprint, or "(mixed: ...)" if boards
  disagree
- Notes column carries SPECIAL_NOTES hits (e.g. source-sense 1 W)

Output: tools/_bom_total.xlsx
"""
from __future__ import annotations
import re
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from gen_board_boms import (
    BOARDS,
    OUT_DIR,
    CATEGORIES,
    SPECIAL_NOTES,
    parse_netlist,
    collect_dnp_refs,
    ref_to_sheet,
    apply_special_note,
    ref_num,
)


def categorize(ref: str, value: str, desc: str) -> str:
    for cat_name, test in CATEGORIES:
        if test(ref, value, desc):
            return cat_name
    return "Other"


def write_total_bom(all_rows: list):
    """all_rows: list of dicts {board, ref, value, footprint, description, sheet}."""
    wb = Workbook()
    first = True

    # Bucket by category (first match wins)
    by_cat = {cat_name: [] for cat_name, _ in CATEGORIES}
    by_cat["Other"] = []
    for c in all_rows:
        cat = categorize(c["ref"], c["value"], c["description"])
        by_cat.setdefault(cat, []).append(c)

    for cat_name, _ in CATEGORIES + [("Other", None)]:
        cat_comps = by_cat.get(cat_name, [])
        if not cat_comps:
            continue
        ws = wb.create_sheet(title=cat_name[:30]) if not first else wb.active
        if first:
            ws.title = cat_name[:30]
            first = False

        # Group by value
        grouped = defaultdict(list)
        for c in cat_comps:
            grouped[c["value"]].append(c)

        headers = ["Value", "Qty", "Boards", "References", "Footprint", "Notes"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDDDDD")
            cell.alignment = Alignment(horizontal="center")

        for value in sorted(grouped.keys()):
            comps = sorted(
                grouped[value],
                key=lambda c: (c["board"], c["ref"][:2], ref_num(c["ref"])),
            )
            refs = ", ".join(f'{c["board"]}:{c["ref"]}' for c in comps)
            boards = sorted(set(c["board"] for c in comps))
            board_display = ", ".join(boards)
            fps = sorted(set(c["footprint"] for c in comps))
            fp_display = fps[0] if len(fps) == 1 else "(mixed: " + " | ".join(fps) + ")"
            notes = sorted(set(filter(
                None,
                (apply_special_note(c["board"], c["ref"]) for c in comps),
            )))
            note_display = "; ".join(notes)
            qty = len(comps)
            ws.append([value, qty, board_display, refs, fp_display, note_display])

        # Totals row at bottom of each tab
        total_qty = sum(len(v) for v in grouped.values())
        ws.append([])
        total_row = ["TOTAL", total_qty, "", "", "", ""]
        ws.append(total_row)
        bold = Font(bold=True)
        for cell in ws[ws.max_row]:
            cell.font = bold

        for col, width in zip("ABCDEF", (18, 6, 24, 50, 50, 50)):
            ws.column_dimensions[col].width = width

    out = OUT_DIR / "_bom_total.xlsx"
    wb.save(out)
    return out


def main():
    all_rows = []
    per_board_counts = {}
    for board, netfile in BOARDS:
        if not Path(netfile).exists():
            print(f"[SKIP] {board}: {netfile} not found")
            continue
        comps = parse_netlist(netfile)
        dnp = collect_dnp_refs(board)
        sheet_map = ref_to_sheet(board)
        kept = 0
        for c in comps:
            if c["ref"] in dnp or c["ref"].startswith("#"):
                continue
            all_rows.append({
                "board": board,
                "ref": c["ref"],
                "value": c["value"],
                "footprint": c["footprint"],
                "description": c["description"],
                "sheet": sheet_map.get(c["ref"], "?"),
            })
            kept += 1
        per_board_counts[board] = (len(comps), kept, len(dnp))
        print(f"=== {board}: netlist={len(comps)}, kept={kept}, DNP={len(dnp)}")

    out = write_total_bom(all_rows)
    print(f"\nTotal consolidated BOM: {out.name} ({len(all_rows)} line items across boards)")
    print("Per-board breakdown (netlist / kept / DNP):")
    for b, (n, k, d) in per_board_counts.items():
        print(f"  {b:12s} {n:4d} / {k:4d} / {d:3d}")


if __name__ == "__main__":
    main()
